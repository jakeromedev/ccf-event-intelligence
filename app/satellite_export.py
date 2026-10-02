"""Single-worksheet export of the active Event's Satellite roster."""

import json
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .satellite_analytics import EFFECTIVE_ASSOCIATIONS_CTE


def satellite_roster_rows(db, event_id, batch_id):
    """Return one row per curated person/effective Satellite, including unassigned."""
    return db.execute(
        EFFECTIVE_ASSOCIATIONS_CTE
        + """
        , representative AS (
            SELECT curated_registrant_id, MIN(registrant_id) registrant_id
            FROM curated_registrant_sources
            WHERE event_id = ? AND batch_id = ?
            GROUP BY curated_registrant_id
        )
        SELECT DISTINCT curated.id curated_id,
               directory.id directory_id,
               hub.id hub_id, hub_group.id group_id,
               COALESCE(directory.name, imported.name, 'Unassigned') satellite,
               COALESCE(hub.name, 'Unassigned') hub,
               COALESCE(hub_group.name, 'Unassigned') hub_group,
               raw.first_name, COALESCE(raw.last_name, curated.last_name) last_name,
               raw.registration_code, raw.source_id,
               raw.b1g_satellite_hub_raw submitted_hub,
               curated.birth_date, curated.birth_month, curated.birth_year,
               raw.source_data_json,
               CASE WHEN hub_group.id IS NOT NULL THEN 'Linked'
                    ELSE 'Needs Mapping' END link_status
        FROM curated_registrants curated
        LEFT JOIN representative rep ON rep.curated_registrant_id = curated.id
        LEFT JOIN registrants raw ON raw.id = rep.registrant_id
        LEFT JOIN effective_associations association
          ON association.curated_registrant_id = curated.id
         AND association.event_id = curated.event_id
         AND association.batch_id = curated.batch_id
        LEFT JOIN satellites imported ON imported.id = association.satellite_id
        LEFT JOIN satellite_directory directory ON directory.id = association.directory_id
        LEFT JOIN satellite_hubs hub ON hub.id = directory.hub_id
        LEFT JOIN hub_groups hub_group ON hub_group.id = hub.hub_group_id
        WHERE curated.event_id = ? AND curated.batch_id = ?
        ORDER BY satellite, hub, hub_group, last_name, first_name, curated_id
        """,
        (event_id, batch_id, event_id, batch_id),
    ).fetchall()


def participant_contact_details(record):
    """Use the same contact and available birth details in the table and export."""
    row = dict(record)
    try:
        source = json.loads(row.get("source_data_json") or "{}")
    except (TypeError, ValueError):
        source = {}
    source = source if isinstance(source, dict) else {}
    return {
        "email": str(source.get("Email Address") or "").strip(),
        "mobile_number": str(source.get("Mobile Number") or "").strip(),
        "date_of_birth": row.get("birth_date") or " ".join(
            str(value) for value in (row.get("birth_month"), row.get("birth_year")) if value
        ),
    }


def satellite_roster_workbook(rows):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Participants by Satellite"
    headers = ("Satellite", "Hub", "Hub Group", "First Name", "Last Name",
               "Date of Birth", "Email", "Contact Number")
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
    unlinked_fill = PatternFill("solid", fgColor="FEE2E2")
    unlinked_font = Font(color="991B1B")
    for row_number, row in enumerate(rows, start=2):
        details = participant_contact_details(row)
        values = (
            row["satellite"], row["hub"], row["hub_group"],
            row["first_name"] or "", row["last_name"] or "",
            details["date_of_birth"], details["email"], details["mobile_number"],
        )
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", str(value)) for value in values])
        # Imported values are always literal text, including IDs and names beginning '='.
        for column in range(1, len(headers) + 1):
            cell = sheet.cell(row_number, column)
            cell.data_type = "s"
            if row["link_status"] != "Linked":
                cell.fill = unlinked_fill
                cell.font = unlinked_font
    sheet.freeze_panes = "D2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate((32, 28, 32, 28, 28, 22, 38, 24), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
