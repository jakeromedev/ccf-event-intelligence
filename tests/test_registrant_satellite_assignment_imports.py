import tempfile
import unittest
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

from app import create_app
from app.aggregation import curated_registrant_detail, satellite_registrants
from app.db import get_db, get_engine
from app.importer import process_batch, store_validation, validate_batch
from app.models import Base
from app.registrant_satellite_assignments import set_manual_satellite_assignment
from app.satellite_analytics import EFFECTIVE_ASSOCIATIONS_CTE, canonical_satellite_metrics
from app.satellite_settings_registrants import event_settings_registrants
from app.satellite_sync import (
    MANUAL_PROTECTED,
    analyze_event_satellite_sync,
    execute_event_satellite_sync,
)
from tests.test_phase1 import (
    BUYER_FIELDS,
    REGISTRANT_REGIONAL_B1G_FIELDS,
    TICKET_FIELDS,
    write_csv,
)


class RegistrantSatelliteAssignmentImportTests(unittest.TestCase):
    def test_effective_assignment_cte_avoids_mysql_reserved_manual_alias(self):
        self.assertNotIn("manual.directory_id", EFFECTIVE_ASSOCIATIONS_CTE)
        self.assertNotIn("manual\n", EFFECTIVE_ASSOCIATIONS_CTE)
        self.assertIn("manual_override", EFFECTIVE_ASSOCIATIONS_CTE)

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.paths = {
            "tickets": root / "tickets.csv",
            "buyers": root / "buyers.csv",
            "registrants": root / "registrants.csv",
        }
        self.app = create_app(
            {
                "TESTING": True,
                "SECRET_KEY": "test",
                "DATABASE_URL": "sqlite+pysqlite:///{}".format(root / "test.sqlite3"),
                "STAGING_DIR": str(root / "staging"),
                "AUTHENTICATION_DISABLED": True,
                "WTF_CSRF_ENABLED": False,
            }
        )
        with self.app.app_context():
            Base.metadata.create_all(get_engine())
            db = get_db()
            self.event_id = db.execute(
                "INSERT INTO events (name) VALUES ('Protected Imports')"
            ).lastrowid
            self.user_id = db.execute(
                """
                INSERT INTO users (
                    username, password_hash, role, status, approved_at
                ) VALUES (
                    'satellite-operator', 'unused-test-hash', 'user',
                    'approved', CURRENT_TIMESTAMP
                )
                """
            ).lastrowid
            group_id = db.execute(
                """
                INSERT INTO hub_groups (code, name, sort_order)
                VALUES ('outside_metro_manila', 'Outside Metro Manila Hubs', 1)
                """
            ).lastrowid
            hub_id = db.execute(
                """
                INSERT INTO satellite_hubs (hub_group_id, name, normalized_name)
                VALUES (?, 'Mindanao South', 'mindanao south')
                """,
                (group_id,),
            ).lastrowid
            self.directory_ids = {}
            for name in ("B1G Tagum", "B1G Davao", "B1G General Santos"):
                self.directory_ids[name] = db.execute(
                    """
                    INSERT INTO satellite_directory (hub_id, name, normalized_name)
                    VALUES (?, ?, ?)
                    """,
                    (hub_id, name, name.casefold()),
                ).lastrowid
            db.commit()

    def tearDown(self):
        self.temp.cleanup()

    def _process(self, imported_satellite, source_id):
        write_csv(
            self.paths["buyers"],
            BUYER_FIELDS,
            [
                {
                    "Id": "buyer-{}".format(source_id),
                    "Slug": "protected-imports",
                    "Event Name": "Protected Imports",
                    "Buyer Reference Number": "BUYER-1",
                    "Payment Status": "Payment Validated",
                    "Quantity": "1",
                }
            ],
        )
        write_csv(
            self.paths["tickets"],
            TICKET_FIELDS,
            [
                {
                    "Id": "ticket-{}".format(source_id),
                    "Slug": "protected-imports",
                    "Event Name": "Protected Imports",
                    "Ticket Code": "T-STABLE",
                    "Control Number": "CONTROL-1",
                    "Ticket Status": "Assigned",
                    "Payment Status": "Payment Validated",
                    "Buyer Reference Number": "BUYER-1",
                }
            ],
        )
        write_csv(
            self.paths["registrants"],
            REGISTRANT_REGIONAL_B1G_FIELDS,
            [
                {
                    "ID": source_id,
                    "Event Name": "Protected Imports",
                    "Event Slug": "protected-imports",
                    "Registration Code": "R-STABLE",
                    "Ticket Code": "T-STABLE",
                    "Ticket Status": "Assigned",
                    "First Name": "Stable",
                    "Last Name": "Registrant",
                    "Email Address": "stable@example.com",
                    "Mobile Number": "09000000000",
                    "Gender": "Female",
                    "Birth Month": "January",
                    "Birth Year": "1990",
                    "Bg Satellite Hub": "Mindanao South",
                    "Mindanao South Hub": imported_satellite,
                }
            ],
        )
        staged = {
            export_type: (str(path), path.name)
            for export_type, path in self.paths.items()
        }
        validation = validate_batch(staged)
        self.assertTrue(validation.valid)
        with self.app.app_context():
            batch_id = store_validation(get_db(), validation, self.event_id)
            process_batch(get_db(), batch_id)
        return batch_id

    def _assignment_values(self, participant_id):
        with self.app.app_context():
            row = get_db().execute(
                """
                SELECT id, event_id, attestation_participant_id, directory_id,
                       assignment_source, source_batch_id, updated_by_user_id,
                       created_at, updated_at
                FROM event_registrant_satellites
                WHERE event_id = ? AND attestation_participant_id = ?
                """,
                (self.event_id, participant_id),
            ).fetchone()
            return tuple(row[key] for key in row.keys())

    def test_satellite_export_uses_active_batch_and_manual_assignment(self):
        self._process("B1G Tagum", "old-source")
        batch_id = self._process("B1G Tagum", "new-source")
        with self.app.app_context():
            db = get_db()
            participant_id = db.execute(
                "SELECT attestation_participant_id FROM attestation_participant_registrants "
                "WHERE event_id = ? AND batch_id = ?",
                (self.event_id, batch_id),
            ).fetchone()[0]
            set_manual_satellite_assignment(
                db, self.event_id, participant_id, self.directory_ids["B1G Davao"]
            )
            db.execute(
                "UPDATE registrants SET first_name = '=1+1', registration_code = '000123' "
                "WHERE batch_id = ?", (batch_id,),
            )
            db.commit()
        client = self.app.test_client()
        response = client.get(
            "/events/{}/satellites/registrants/export.xlsx?q=missing&satellite=-1&page=99".format(
                self.event_id
            )
        )
        self.assertEqual(200, response.status_code)
        self.assertIn("attachment;", response.headers["Content-Disposition"])
        self.assertEqual("private, no-store", response.headers["Cache-Control"])
        workbook = load_workbook(BytesIO(response.data))
        self.assertEqual(["Participants by Satellite"], workbook.sheetnames)
        sheet = workbook.active
        self.assertEqual(2, sheet.max_row)
        self.assertEqual("B1G Davao", sheet["A2"].value)
        self.assertEqual("=1+1", sheet["D2"].value)
        self.assertEqual("s", sheet["D2"].data_type)
        self.assertEqual("Registrant", sheet["E2"].value)
        self.assertTrue(sheet["F2"].value)
        self.assertEqual("stable@example.com", sheet["G2"].value)
        self.assertEqual("09000000000", sheet["H2"].value)
        self.assertEqual("s", sheet["H2"].data_type)
        self.assertEqual("A1:H2", sheet.auto_filter.ref)
        self.assertEqual(
            ["Satellite", "Hub", "Hub Group", "First Name", "Last Name", "Date of Birth", "Email", "Contact Number"],
            [cell.value for cell in sheet[1]],
        )
        self.assertIsNone(sheet["A2"].fill.patternType)
        page = client.get("/events/{}/satellites".format(self.event_id))
        self.assertIn(b"Download all participants (.xlsx)", page.data)
        self.assertNotIn(b'class="satellite-group-summary"', page.data)
        key = "directory:{}".format(self.directory_ids["B1G Davao"])
        page = client.get("/events/{}/satellites".format(self.event_id), query_string={
            "roster_satellite": key, "roster_q": "000123",
        })
        self.assertEqual(200, page.status_code)
        table = page.data.split(b'id="satellite-ranking-table"')[1].split(b'</section>')[0]
        self.assertIn(b"<td>=1+1</td><td>Registrant</td>", table)
        self.assertNotIn(b"<th>Registration ID</th>", table)
        self.assertNotIn(b"<th>Link Status</th>", table)
        self.assertIn(b"stable@example.com", table)
        self.assertIn(b"09000000000", table)
        self.assertIn(b"<th>Date of Birth</th>", table)
        self.assertNotIn(b'class="satellite-participant-unlinked"', table)
        self.assertIn(('value="{}" selected'.format(key)).encode(), table)
        self.assertIn(b"Clear search", table)
        page = client.get("/events/{}/satellites".format(self.event_id), query_string={
            "roster_satellite": key, "roster_q": "no match",
        })
        self.assertIn(b"No participants match the selected filters and search.", page.data)

    def test_satellite_export_includes_unmapped_and_unassigned_registrants(self):
        batch_id = self._process("B1G Tagum", "source-1")
        with self.app.app_context():
            db = get_db()
            db.execute("UPDATE satellites SET directory_id = NULL WHERE batch_id = ?", (batch_id,))
            db.commit()
        client = self.app.test_client()
        url = "/events/{}/satellites/registrants/export.xlsx".format(self.event_id)
        sheet = load_workbook(BytesIO(client.get(url).data)).active
        self.assertEqual("B1G Tagum", sheet["A2"].value)
        for cell in sheet[2]:
            self.assertEqual("solid", cell.fill.patternType)
            self.assertEqual("00FEE2E2", cell.fill.fgColor.rgb)
        page = client.get("/events/{}/satellites".format(self.event_id))
        self.assertIn(b'class="satellite-participant-unlinked"', page.data)
        with self.app.app_context():
            db = get_db()
            db.execute("DELETE FROM curated_registrant_satellites WHERE batch_id = ?", (batch_id,))
            db.commit()
        sheet = load_workbook(BytesIO(client.get(url).data)).active
        self.assertEqual(2, sheet.max_row)
        self.assertEqual("Unassigned", sheet["A2"].value)
        self.assertEqual("00FEE2E2", sheet["H2"].fill.fgColor.rgb)

    def test_satellite_export_requires_event_batch_and_authentication(self):
        client = self.app.test_client()
        url = "/events/{}/satellites/registrants/export.xlsx".format(self.event_id)
        self.assertEqual(404, client.get(url).status_code)
        self.assertEqual(404, client.get("/events/999999/satellites/registrants/export.xlsx").status_code)
        self.app.config["AUTHENTICATION_DISABLED"] = False
        self.assertEqual(302, client.get(url).status_code)

    def test_satellite_export_contains_entire_roster_grouped_by_satellite(self):
        batch_id = self._process("B1G Tagum", "source-1")
        with self.app.app_context():
            db = get_db()
            satellite_id = db.execute(
                "SELECT id FROM satellites WHERE batch_id = ?", (batch_id,)
            ).fetchone()[0]
            for index in range(105):
                curated_id = db.execute(
                    "INSERT INTO curated_registrants "
                    "(event_id, batch_id, last_name, dedupe_key, dedupe_status, registration_type) "
                    "VALUES (?, ?, ?, ?, 'incomplete', 'participant')",
                    (self.event_id, batch_id, "Person {:03}".format(index), "extra-{}".format(index)),
                ).lastrowid
                if index % 2:
                    db.execute(
                        "INSERT INTO curated_registrant_satellites "
                        "(event_id, batch_id, curated_registrant_id, satellite_id) VALUES (?, ?, ?, ?)",
                        (self.event_id, batch_id, curated_id, satellite_id),
                    )
            db.commit()
        response = self.app.test_client().get(
            "/events/{}/satellites/registrants/export.xlsx".format(self.event_id)
        )
        sheet = load_workbook(BytesIO(response.data)).active
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        self.assertEqual(106, len(rows))
        satellites = [row[0] for row in rows]
        self.assertEqual(sorted(satellites), satellites)
        self.assertEqual(53, satellites.count("B1G Tagum"))
        self.assertEqual(53, satellites.count("Unassigned"))

    def test_manual_assignment_survives_same_different_and_multiple_future_batches(self):
        first_batch_id = self._process("B1G Tagum", "source-1")
        with self.app.app_context():
            db = get_db()
            participant_id = db.execute(
                """
                SELECT attestation_participant_id
                FROM attestation_participant_registrants
                WHERE event_id = ? AND batch_id = ?
                """,
                (self.event_id, first_batch_id),
            ).fetchone()["attestation_participant_id"]
            set_manual_satellite_assignment(
                db,
                self.event_id,
                participant_id,
                self.directory_ids["B1G Davao"],
                updated_by_user_id=self.user_id,
            )
            db.commit()
        protected_values = self._assignment_values(participant_id)

        with self.app.app_context():
            db = get_db()
            sync_plan = analyze_event_satellite_sync(db, self.event_id)
            sync_result = execute_event_satellite_sync(db, self.event_id)
            db.commit()
        self.assertEqual(MANUAL_PROTECTED, sync_plan["registrations"][0]["status"])
        self.assertEqual(0, sync_result["synchronized_count"])
        self.assertEqual(0, sync_result["not_synced_count"])
        self.assertEqual(protected_values, self._assignment_values(participant_id))

        with self.app.app_context():
            db = get_db()
            metrics = canonical_satellite_metrics(db, first_batch_id)
            davao_roster = satellite_registrants(
                db, first_batch_id, self.directory_ids["B1G Davao"]
            )
            tagum_roster = satellite_registrants(
                db, first_batch_id, self.directory_ids["B1G Tagum"]
            )
            curated_id = db.execute(
                "SELECT id FROM curated_registrants WHERE batch_id = ?",
                (first_batch_id,),
            ).fetchone()["id"]
            curated = curated_registrant_detail(db, first_batch_id, curated_id)
        ranking = {item["name"]: item["registrants"] for item in metrics["satellites"]}
        self.assertEqual(1, ranking["B1G Davao"])
        self.assertNotIn("B1G Tagum", ranking)
        self.assertEqual(1, davao_roster["registrants"])
        self.assertIsNone(tagum_roster)
        self.assertEqual("B1G Davao", curated["effective_satellites"][0]["name"])
        self.assertEqual("manual", curated["effective_satellites"][0]["assignment_source"])

        expected_imports = (
            ("B1G Tagum", "source-2"),
            ("B1G General Santos", "source-3"),
            ("B1G Tagum", "source-4"),
        )
        for imported_satellite, source_id in expected_imports:
            with self.subTest(imported_satellite=imported_satellite, source_id=source_id):
                batch_id = self._process(imported_satellite, source_id)
                self.assertEqual(protected_values, self._assignment_values(participant_id))
                with self.app.app_context():
                    db = get_db()
                    current_participant_id = db.execute(
                        """
                        SELECT attestation_participant_id
                        FROM attestation_participant_registrants
                        WHERE event_id = ? AND batch_id = ?
                        """,
                        (self.event_id, batch_id),
                    ).fetchone()["attestation_participant_id"]
                    payload = event_settings_registrants(db, self.event_id)
                    assignment_count = db.execute(
                        """
                        SELECT COUNT(*) count FROM event_registrant_satellites
                        WHERE event_id = ? AND attestation_participant_id = ?
                        """,
                        (self.event_id, participant_id),
                    ).fetchone()["count"]

                self.assertEqual(participant_id, current_participant_id)
                self.assertEqual(1, assignment_count)
                self.assertEqual(imported_satellite, payload["rows"][0]["imported_satellite"])
                self.assertEqual("B1G Davao", payload["rows"][0]["effective_satellite"])
                self.assertEqual("manual", payload["rows"][0]["assignment_source"])
                self.assertEqual(
                    "satellite-operator", payload["rows"][0]["assignment_updated_by"]
                )

        with self.app.app_context():
            evidence = get_db().execute(
                """
                SELECT record.source_id, record.satellite_name, batch.status
                FROM registrants record
                JOIN import_batches batch ON batch.id = record.batch_id
                WHERE batch.event_id = ?
                ORDER BY batch.id
                """,
                (self.event_id,),
            ).fetchall()
        self.assertEqual(
            ["source-1", "source-2", "source-3", "source-4"],
            [row["source_id"] for row in evidence],
        )
        self.assertEqual(
            ["B1G Tagum", "B1G Tagum", "B1G General Santos", "B1G Tagum"],
            [row["satellite_name"] for row in evidence],
        )
        self.assertEqual(1, sum(row["status"] == "active" for row in evidence))


if __name__ == "__main__":
    unittest.main()
